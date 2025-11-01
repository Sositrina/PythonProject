import csv
import json
import os
from typing import Any, Dict, List

import openpyxl

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DATA_DIR = os.path.join(os.path.dirname(BASE_DIR), "data")
TRANSACTIONS_DIR = os.path.join(os.path.dirname(BASE_DIR), "transactions")


def read_json_file(file_name: str) -> List[Dict[str, Any]]:
    file_path = os.path.join(DATA_DIR, file_name)
    print(f"Ищу JSON файл: {file_path}")
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"Файл {file_path} не найден")

    with open(file_path, "r", encoding="utf-8") as f:
        data: Any = json.load(f)
        if isinstance(data, list):
            return data  # type: ignore[return-value]
        raise ValueError("JSON-файл должен содержать список транзакций")


def read_csv_file(file_name: str) -> List[Dict[str, Any]]:
    file_path = os.path.join(TRANSACTIONS_DIR, file_name)
    print(f"Ищу CSV файл: {file_path}")
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"Файл {file_path} не найден")

    transactions: List[Dict[str, Any]] = []
    with open(file_path, newline="", encoding="utf-8") as f:
        reader = csv.DictReader(f, delimiter=";")
        for row in reader:
            transactions.append(
                {
                    "date": row.get("date") or "",
                    "description": row.get("description") or "",
                    "from": row.get("from") or "",
                    "to": row.get("to") or "",
                    "amount": row.get("amount") or "",
                    "currency": {"name": row.get("currency_name") or "", "code": row.get("currency_code") or ""},
                    "state": (row.get("state") or "").upper(),
                }
            )
    return transactions


def read_xlsx_file(file_name: str) -> List[Dict[str, Any]]:
    file_path = os.path.join(TRANSACTIONS_DIR, file_name)
    print(f"Ищу XLSX файл: {file_path}")
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"Файл {file_path} не найден")
    wb = openpyxl.load_workbook(file_path)
    sheet = wb.active
    transactions: List[Dict[str, Any]] = []
    headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    for row in sheet.iter_rows(min_row=2, values_only=True):
        transaction = dict(zip(headers, row))
        state_value = transaction.get("state")
        transactions.append(
            {
                "date": transaction.get("date") or "",
                "description": transaction.get("description") or "",
                "from": transaction.get("from") or "",
                "to": transaction.get("to") or "",
                "amount": transaction.get("amount") or "",
                "currency": {
                    "name": transaction.get("currency_name") or "",
                    "code": transaction.get("currency_code") or "",
                },
                "state": state_value.upper() if state_value else "",
            }
        )
    return transactions


def get_transactions(choice: int) -> List[Dict[str, Any]]:
    if choice == 1:
        return read_json_file("operations.json")
    elif choice == 2:
        return read_csv_file("transactions.csv")
    elif choice == 3:
        return read_xlsx_file("transactions_excel.xlsx")
